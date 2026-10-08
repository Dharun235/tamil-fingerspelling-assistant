#!/usr/bin/env python3
"""Create the editable judge/demo presentation deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs/assets"
OUT = ROOT / "docs/Tamil_Fingerspelling_Demo_Deck.pptx"

NAVY = RGBColor(18, 59, 93)
TEAL = RGBColor(43, 120, 160)
INK = RGBColor(28, 43, 52)
MUTED = RGBColor(75, 96, 108)
PALE = RGBColor(245, 249, 251)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(210, 155, 58)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Arial"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def image(slide, path, x, y, w=None, h=None, alt=None):
    kwargs = {}
    if w is not None:
        kwargs["width"] = Inches(w)
    if h is not None:
        kwargs["height"] = Inches(h)
    shape = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)
    if alt:
        shape._element.nvPicPr.cNvPr.set("descr", alt)
    return shape


def base(title, kicker=None):
    slide = prs.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, PALE)
    rect(slide, 0, 0, 13.333, 0.16, TEAL)
    if kicker:
        text(slide, kicker.upper(), 0.55, 0.38, 5.0, 0.25, 10, TEAL, True)
    text(slide, title, 0.55, 0.68, 12.0, 0.52, 27, NAVY, True)
    return slide


def footer(slide, number):
    text(slide, "Real-Time Tamil Fingerspelling Assistant", 0.55, 7.12, 5.5, 0.2, 8.5, MUTED)
    text(slide, str(number), 12.45, 7.12, 0.35, 0.2, 8.5, MUTED, align=PP_ALIGN.RIGHT)


def bullets(slide, items, x, y, w, size=19, color=INK, gap=0.12):
    for item in items:
        text(slide, "• " + item, x, y, w, 0.43, size, color)
        y += 0.43 + gap


# 1 — title
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, NAVY)
rect(slide, 0, 0, 13.333, 0.18, TEAL)
text(slide, "Tamil Fingerspelling\nAssistant", 0.75, 0.8, 5.4, 1.35, 35, WHITE, True)
text(slide, "Bi-directional, real-time Tamil fingerspelling assistance", 0.8, 2.45, 5.3, 0.55, 19, RGBColor(205, 229, 239))
text(slide, "OpenCV 5  •  AWS  •  Accessibility", 0.8, 3.25, 5.2, 0.35, 14, RGBColor(158, 202, 220), True)
rect(slide, 6.35, 0.95, 5.9, 5.15, RGBColor(29, 78, 111), True, RGBColor(73, 145, 176))
text(slide, "One assistant.\nTwo directions.", 6.8, 1.35, 5.0, 0.95, 31, WHITE, True, PP_ALIGN.CENTER)
rect(slide, 6.85, 2.75, 4.9, 0.75, RGBColor(235, 247, 251), True, RGBColor(112, 184, 207))
text(slide, "Fingerspelling  →  Tamil text", 7.1, 2.94, 4.4, 0.3, 19, NAVY, True, PP_ALIGN.CENTER)
rect(slide, 6.85, 3.85, 4.9, 0.75, RGBColor(255, 244, 224), True, RGBColor(220, 174, 94))
text(slide, "Tamil text  →  visual sign", 7.1, 4.04, 4.4, 0.3, 19, NAVY, True, PP_ALIGN.CENTER)
text(slide, "OpenCV 5  •  AWS  •  Accessibility", 6.8, 5.25, 5.0, 0.3, 14, RGBColor(205, 229, 239), True, PP_ALIGN.CENTER)
rect(slide, 0.8, 5.75, 5.15, 0.72, RGBColor(29, 78, 111), True, RGBColor(73, 145, 176))
text(slide, "Turn hand poses into Tamil text —\nand text back into visual signs.", 1.05, 5.87, 4.7, 0.45, 16, WHITE, True)

# 2 — motivation
slide = base("Why this matters", "01 · Motivation")
rect(slide, 0.65, 1.55, 5.3, 4.95, WHITE, True, RGBColor(215, 226, 232))
text(slide, "Tamil fingerspelling needs\nlanguage-specific support", 1.0, 1.9, 4.55, 0.85, 25, NAVY, True)
bullets(slide, [
    "Tamil has 247 written characters represented in TLFS23.",
    "Vernacular fingerspelling resources remain limited.",
    "Users need feedback before a character is committed.",
    "Reference signs help learners and communication partners.",
], 1.0, 3.05, 4.45, 18, INK, 0.2)
text(slide, "Not continuous sign-language translation.\nAn assistive, human-in-the-loop prototype.", 1.0, 5.65, 4.45, 0.5, 14, TEAL, True)
image(slide, ASSETS / "tlfs23-sign-mapping.png", 6.3, 1.65, w=6.4, alt="Representative TLFS23 class IDs, Tamil labels, and reference sign images.")
text(slide, "TLFS23 provides class labels and reference signs.", 6.55, 6.22, 6.0, 0.28, 12, MUTED, align=PP_ALIGN.CENTER)
footer(slide, 2)

# 3 — capabilities
slide = base("What we built", "02 · Product")
cards = [
    (0.75, "Recognize", "Webcam frames → hand landmarks → Tamil character", "#e8f3f8"),
    (4.55, "Guide", "Green/red finger feedback + stable-hold commitment", "#fff3df"),
    (8.35, "Reverse", "Typed Tamil → TLFS23 visual reference sign", "#eaf4ee"),
]
for x, title, subtitle, fill in cards:
    rect(slide, x, 1.7, 3.45, 3.05, RGBColor.from_string(fill[1:]), True, RGBColor(190, 207, 216))
    text(slide, title, x + 0.3, 2.15, 2.85, 0.42, 24, NAVY, True, PP_ALIGN.CENTER)
    text(slide, subtitle, x + 0.35, 2.95, 2.75, 1.0, 18, INK, align=PP_ALIGN.CENTER)
text(slide, "Human-in-the-loop design: users can correct framing before the system commits a character.", 1.1, 5.45, 11.1, 0.55, 20, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 3)

# 4 — technical flow
slide = base("How one frame becomes Tamil text", "03 · OpenCV 5 pipeline")
image(slide, ASSETS / "recognition-flow.png", 0.65, 1.45, w=12.05, alt="Recognition flow from browser webcam through FastAPI, OpenCV 5, hand pose, geometry, mapping, stability, and Tamil UI.")
text(slide, "OpenCV 5 decodes and normalizes every frame before hand-pose inference.", 1.0, 6.58, 11.4, 0.3, 14, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 4)

# 5 — AWS
slide = base("Reproducible AWS deployment", "04 · Cloud delivery")
image(slide, ASSETS / "aws-deployment-flow.png", 0.65, 1.45, w=12.05, alt="AWS deployment flow from source and ECR through ECS Express Mode and Fargate to the public endpoint, CloudWatch, and browser user.")
text(slide, "Live endpoint: https://ta-89507d6158f8458c88ac7b38194d8efc.ecs.eu-north-1.on.aws", 0.8, 6.58, 11.8, 0.3, 12, TEAL, True, PP_ALIGN.CENTER)
footer(slide, 5)

# 6 — evidence and close
slide = base("Evidence, limits, and next steps", "05 · Results")
rect(slide, 0.65, 1.45, 5.85, 4.95, WHITE, True, RGBColor(215, 226, 232))
text(slide, "Measured evidence", 1.0, 1.83, 4.7, 0.35, 22, NAVY, True)
bullets(slide, [
    "4,940 TLFS23 images across 247 classes.",
    "92.5–97.3% finger-state accuracy.",
    "Mac local inference: 116.6 ms mean.",
    "AWS warm WebSocket: 113.8 ms mean.",
    "User commit includes 350 ms stability hold.",
], 1.0, 2.38, 4.85, 18, INK, 0.15)
text(slide, "Limitations", 1.0, 5.15, 4.7, 0.3, 18, TEAL, True)
text(slide, "Occlusion, framing, lighting, and merged hands can reduce accuracy.", 1.0, 5.52, 4.8, 0.5, 14, INK)
rect(slide, 6.85, 1.45, 5.8, 4.95, NAVY, True, TEAL)
text(slide, "Next", 7.3, 1.83, 4.8, 0.35, 22, WHITE, True)
bullets(slide, [
    "Measure full end-to-end user commit latency.",
    "Improve two-hand and occlusion handling.",
    "Expand real-world evaluation across users.",
    "Explore speech output and more Tamil sequences.",
], 7.3, 2.45, 4.7, 18, RGBColor(231, 242, 247), 0.2)
text(slide, "A browser-based assistant turns Tamil fingerspelling into visible text while helping users correct their pose in real time.", 7.3, 5.25, 4.75, 0.75, 17, WHITE, True)
footer(slide, 6)

prs.save(OUT)
print(OUT)
